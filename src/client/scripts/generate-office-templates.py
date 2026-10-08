#!/usr/bin/env python3
"""Regenerate the blank Office templates behind the sidebar's
"New Word Document / New Excel Spreadsheet / New PowerPoint Presentation" items.

The output is committed; rerun only when the templates need to change.
python-docx, openpyxl and python-pptx are MIT/BSD licensed, so nothing here is
derived from ONLYOFFICE's AGPL templates.

Usage (from src/client):
    python3 -m venv "${TMPDIR:-/tmp}/guideants-office-templates-venv"
    "${TMPDIR:-/tmp}/guideants-office-templates-venv/bin/pip" install \
        "python-docx>=1.1,<2" "openpyxl>=3.1,<4" "python-pptx>=1.0,<2"
    "${TMPDIR:-/tmp}/guideants-office-templates-venv/bin/python" scripts/generate-office-templates.py
"""
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation

OUT_DIR = Path(__file__).resolve().parent.parent / "src" / "assets" / "office-templates"

# python-docx / python-pptx credit themselves in docProps/core.xml; a user's new
# document shouldn't.
CORE_TEXT_FIELDS = ("author", "last_modified_by", "title", "subject", "keywords", "comments", "category")


def scrub_core_properties(props) -> None:
    for field in CORE_TEXT_FIELDS:
        setattr(props, field, "")


def build_docx(path: Path) -> None:
    doc = Document()
    if not doc.paragraphs:
        doc.add_paragraph()
    scrub_core_properties(doc.core_properties)
    doc.save(path)


def build_xlsx(path: Path) -> None:
    wb = Workbook()
    wb.active.title = "Sheet1"
    wb.properties.creator = None
    wb.properties.lastModifiedBy = None
    wb.save(path)


def build_pptx(path: Path) -> None:
    # python-pptx's default master is laid out for 4:3, so keep that size; users can
    # change it in the editor. "Title Slide" (layout 0) gives the usual empty
    # click-to-add placeholders instead of a zero-slide deck.
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[0])
    scrub_core_properties(prs.core_properties)
    prs.save(path)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    build_docx(OUT_DIR / "blank.docx")
    build_xlsx(OUT_DIR / "blank.xlsx")
    build_pptx(OUT_DIR / "blank.pptx")
    for name in ("blank.docx", "blank.xlsx", "blank.pptx"):
        print(f"wrote {OUT_DIR / name} ({(OUT_DIR / name).stat().st_size} bytes)")


if __name__ == "__main__":
    main()
